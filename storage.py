"""Persistência do R.A - Equivalencias em uma pasta de trabalho Excel."""

from __future__ import annotations

from pathlib import Path
from typing import Any
from datetime import datetime
import hashlib
import hmac
import os
import re


PROCESS_COLUMNS = [
    ("ID_Processo", "id_processo"),
    ("Data_Requerimento", "data_requerimento"),
    ("ID_Aluno", "id_aluno"),
    ("Nome_Aluno", "nome_aluno"),
    ("CPF", "cpf"),
    ("RG", "rg"),
    ("Data_Nascimento", "data_nascimento"),
    ("Data_Matricula", "data_matricula"),
    ("Telefone", "telefone"),
    ("Email", "email"),
    ("Situacao_Interessado", "situacao_interessado"),
    ("Curso_Origem", "curso_origem"),
    ("Ano_Cursado", "ano_cursado"),
    ("Tipo_Aluno", "tipo_aluno"),
    ("Curso_Destino", "curso_destino"),
    ("PC_Cursado", "pc_cursado"),
    ("PC_Destino", "pc_destino"),
    ("Codigo_Oferta", "codigo_oferta"),
    ("Periodo_Letivo", "periodo_letivo"),
    ("CH_Total_Oferta", "ch_total_oferta"),
    ("CH_A_Cursar", "ch_a_cursar"),
    ("Valor_Total", "valor_total"),
    ("Qtd_Parcelas", "qtd_parcelas"),
    ("Valor_Parcela", "valor_parcela"),
    ("Status_Processo", "status_processo"),
    ("Responsavel_Financeiro", "responsavel_financeiro"),
    ("CPF_Responsavel", "cpf_responsavel"),
    ("Prazo_Resposta_Dias", "prazo_resposta_dias"),
    ("Observacoes", "observacoes"),
]

EQUIVALENCE_COLUMNS = [
    ("ID_Processo", "id_processo"),
    ("Modulo_UC_Cursado", "modulo_uc_cursado"),
    ("Nome_UC_Cursado", "nome_uc_cursado"),
    ("Mencao", "mencao"),
    ("Numero_UC_Destino", "numero_uc_destino"),
    ("Nome_UC_Destino", "nome_uc_destino"),
    ("Numero_Aula", "numero_aula"),
    ("ID_Curso", "id_curso"),
]

PROCESS_SHEET = "Processos"
EQUIVALENCE_SHEET = "Equivalencias"
USER_SHEET = "Usuarios"
STUDENT_SHEET = "Alunos"

STUDENT_COLUMNS = [
    ("ID_Aluno", "id_aluno"),
    ("Nome_Aluno", "nome_aluno"),
    ("CPF", "cpf"),
    ("RG", "rg"),
    ("Data_Nascimento", "data_nascimento"),
    ("Data_Matricula", "data_matricula"),
    ("Telefone", "telefone"),
    ("Email", "email"),
    ("Curso_Origem", "curso_origem"),
    ("Ano_Cursado", "ano_cursado"),
    ("Importado_Em", "importado_em"),
    ("IDs_Processos", "ids_processos"),
]

USER_COLUMNS = [
    ("Usuario", "username"),
    ("Funcao", "function"),
    ("Salt_PBKDF2", "salt"),
    ("Hash_PBKDF2_SHA256", "password_hash"),
    ("Trocar_Senha_No_Login", "must_change_password"),
]
PASSWORD_HASH_ITERATIONS = 600_000


class ExcelStorageError(Exception):
    """Erro legível para problemas ao abrir ou atualizar a pasta Excel."""


class ExcelStorage:
    def __init__(self, path: str | Path):
        self.path = Path(path).expanduser()

    @staticmethod
    def _load_openpyxl():
        try:
            import openpyxl
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError as exc:
            raise ExcelStorageError(
                "A biblioteca openpyxl não está instalada neste interpretador. "
                "Abra run_sistema.bat para preparar o ambiente local. Se estiver usando o IDE, "
                "selecione .venv\\Scripts\\python.exe como interpretador ou instale a dependência "
                "no ambiente escolhido."
            ) from exc
        return openpyxl, Alignment, Font, PatternFill

    @staticmethod
    def _key(value: Any) -> str:
        return re.sub(r"[^a-z0-9]", "", str(value or "").strip().casefold())

    def _open(self):
        openpyxl, Alignment, Font, PatternFill = self._load_openpyxl()
        self.path.parent.mkdir(parents=True, exist_ok=True)
        if self.path.exists():
            try:
                wb = openpyxl.load_workbook(self.path)
            except PermissionError as exc:
                raise ExcelStorageError(
                    "Não consegui abrir a planilha. Feche o arquivo no Excel e tente novamente."
                ) from exc
            except Exception as exc:
                raise ExcelStorageError(
                    f"Não consegui ler a planilha selecionada:\n{exc}"
                ) from exc
        else:
            wb = openpyxl.Workbook()
            wb.active.title = PROCESS_SHEET

        self._ensure_sheet(wb, PROCESS_SHEET, PROCESS_COLUMNS, Font, PatternFill, Alignment)
        self._ensure_sheet(wb, EQUIVALENCE_SHEET, EQUIVALENCE_COLUMNS, Font, PatternFill, Alignment)
        self._ensure_sheet(wb, USER_SHEET, USER_COLUMNS, Font, PatternFill, Alignment)
        return wb

    @staticmethod
    def _ensure_sheet(wb, title, columns, Font, PatternFill, Alignment):
        if title in wb.sheetnames:
            ws = wb[title]
        elif title == PROCESS_SHEET and len(wb.sheetnames) == 1:
            ws = wb.active
            if ws.max_row == 1 and ws.max_column == 1 and ws.cell(1, 1).value is None:
                ws.title = title
            else:
                ws = wb.create_sheet(title)
        else:
            ws = wb.create_sheet(title)

        existing = {}
        for col_idx in range(1, ws.max_column + 1):
            value = ws.cell(1, col_idx).value
            if value:
                existing[ExcelStorage._key(value)] = col_idx

        for heading, _field in columns:
            key = ExcelStorage._key(heading)
            if key not in existing:
                col_idx = ws.max_column + 1 if ws.max_row > 1 or ws.cell(1, 1).value else 1
                ws.cell(1, col_idx, heading)
                existing[key] = col_idx

        header_fill = PatternFill("solid", fgColor="004A8D")
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(1, col_idx)
            cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 28

        for col_idx, (heading, _field) in enumerate(columns, start=1):
            # Set readable widths for columns added by this application.
            if col_idx <= ws.max_column and ws.cell(1, col_idx).value == heading:
                ws.column_dimensions[ws.cell(1, col_idx).column_letter].width = min(
                    max(len(heading) + 3, 16), 34
                )

    @staticmethod
    def _header_map(ws, columns):
        by_heading = {ExcelStorage._key(heading): field for heading, field in columns}
        mapping = {}
        for col_idx in range(1, ws.max_column + 1):
            header = ws.cell(1, col_idx).value
            if header is None:
                continue
            key = ExcelStorage._key(header)
            if key in by_heading:
                mapping[by_heading[key]] = col_idx
        return mapping

    @staticmethod
    def _safe_save(wb, path: Path):
        temp_path = path.with_name(f".{path.stem}.tmp{path.suffix}")
        try:
            wb.save(temp_path)
            temp_path.replace(path)
        except PermissionError as exc:
            try:
                temp_path.unlink(missing_ok=True)
            except OSError:
                pass
            raise ExcelStorageError(
                "Não consegui salvar. Feche a planilha no Excel e tente novamente."
            ) from exc
        except Exception as exc:
            try:
                temp_path.unlink(missing_ok=True)
            except OSError:
                pass
            raise ExcelStorageError(f"Falha ao salvar a planilha:\n{exc}") from exc

    def initialize(self):
        wb = self._open()
        self._safe_save(wb, self.path)

    def list_processes(self) -> list[dict[str, str]]:
        wb = self._open()
        ws = wb[PROCESS_SHEET]
        mapping = self._header_map(ws, PROCESS_COLUMNS)
        records = []
        for row_idx in range(2, ws.max_row + 1):
            record = {}
            for _heading, field in PROCESS_COLUMNS:
                col_idx = mapping.get(field)
                value = ws.cell(row_idx, col_idx).value if col_idx else None
                record[field] = "" if value is None else str(value)
            if any(record.values()):
                records.append(record)
        return records

    def find_processes(self, query: str) -> list[dict[str, str]]:
        query = query.strip().casefold()
        records = self.list_processes()
        if not query:
            return records
        return [
            record
            for record in records
            if any(query in value.casefold() for value in record.values())
        ]

    def get_process(self, process_id: str) -> dict[str, str] | None:
        target = process_id.strip().casefold()
        for record in self.list_processes():
            if record.get("id_processo", "").strip().casefold() == target:
                return record
        return None

    def next_process_id(self, year: int, prefix: str = "PROC") -> str:
        prefix = re.sub(r"[^A-Za-z0-9]+", "-", str(prefix or "PROC")).strip("-").upper() or "PROC"
        pattern = re.compile(rf"^{re.escape(prefix)}-{year}-(\d+)$", re.IGNORECASE)
        numbers = []
        for record in self.list_processes():
            match = pattern.match(record.get("id_processo", "").strip())
            if match:
                numbers.append(int(match.group(1)))
        return f"{prefix}-{year}-{max(numbers, default=0) + 1:03d}"

    def save_process(self, record: dict[str, str], original_id: str | None = None):
        wb = self._open()
        ws = wb[PROCESS_SHEET]
        mapping = self._header_map(ws, PROCESS_COLUMNS)
        lookup_id = (original_id or record.get("id_processo", "")).strip().casefold()
        row_idx = None
        id_col = mapping.get("id_processo")
        if id_col:
            for candidate in range(2, ws.max_row + 1):
                value = str(ws.cell(candidate, id_col).value or "").strip().casefold()
                if value == lookup_id:
                    row_idx = candidate
                    break

        if row_idx is None:
            row_idx = ws.max_row + 1

        for _heading, field in PROCESS_COLUMNS:
            col_idx = mapping.get(field)
            if col_idx:
                cell = ws.cell(row_idx, col_idx)
                cell.value = str(record.get(field, "") or "")
                cell.data_type = "s"
                cell.number_format = "@"
        self._safe_save(wb, self.path)

    def _remove_process_link(self, wb, student_id: str, process_id: str):
        if not student_id or not process_id:
            return
        ws = wb[STUDENT_SHEET]
        mapping = self._header_map(ws, STUDENT_COLUMNS)
        id_col = mapping.get("id_aluno")
        links_col = mapping.get("ids_processos")
        if not id_col or not links_col:
            return
        for row_idx in range(2, ws.max_row + 1):
            saved_id = str(ws.cell(row_idx, id_col).value or "").strip()
            if saved_id.casefold() != student_id.casefold():
                continue
            links = [
                value.strip()
                for value in str(ws.cell(row_idx, links_col).value or "").split(";")
                if value.strip() and value.strip().casefold() != process_id.casefold()
            ]
            ws.cell(row_idx, links_col).value = "; ".join(links)
            return

    @staticmethod
    def _student_identity(record: dict[str, str]) -> tuple[str, str]:
        student_id = str(record.get("id_aluno", "") or "").strip().casefold()
        cpf = re.sub(r"\D", "", str(record.get("cpf", "") or ""))
        return student_id, cpf

    def _resolve_student_identity(self, ws, mapping, record: dict[str, str]) -> tuple[int | None, str]:
        """Localiza uma única ficha sem permitir que ID e CPF apontem para pessoas diferentes."""
        student_id, cpf = self._student_identity(record)
        id_col = mapping.get("id_aluno")
        cpf_col = mapping.get("cpf")
        identities = {}
        for candidate in range(2, ws.max_row + 1):
            saved_id = str(ws.cell(candidate, id_col).value or "").strip().casefold() if id_col else ""
            saved_cpf = re.sub(r"\D", "", str(ws.cell(candidate, cpf_col).value or "")) if cpf_col else ""
            identities[candidate] = (saved_id, saved_cpf)
        id_rows = {row for row, (saved_id, _) in identities.items() if student_id and saved_id == student_id}
        cpf_rows = {row for row, (_, saved_cpf) in identities.items() if cpf and saved_cpf == cpf}
        matched_rows = id_rows | cpf_rows
        if len(id_rows) > 1 or len(cpf_rows) > 1:
            lines = ", ".join(str(row) for row in sorted(matched_rows))
            raise ExcelStorageError(
                f"A base de alunos contém ID ou CPF duplicado nas linhas {lines}. "
                "Revise esses cadastros antes de salvar."
            )
        if len(matched_rows) > 1:
            raise ExcelStorageError(
                f"Conflito: o ID '{str(record.get('id_aluno', '') or '').strip()}' "
                "e o CPF informado pertencem a alunos diferentes. Revise o ID e o CPF."
            )
        row_idx = next(iter(matched_rows), None)
        if row_idx is None:
            return None, str(record.get("id_aluno", "") or "").strip()
        saved_id, saved_cpf = identities[row_idx]
        if cpf and saved_cpf and cpf != saved_cpf:
            raise ExcelStorageError(
                f"Conflito: o ID '{str(record.get('id_aluno', '') or '').strip()}' "
                f"já está cadastrado com outro CPF (linha {row_idx} da base de alunos). "
                "Revise o ID e o CPF."
            )
        canonical_rows = [
            row for row, (other_id, other_cpf) in identities.items()
            if (saved_id and other_id == saved_id) or (saved_cpf and other_cpf == saved_cpf)
        ]
        if len(canonical_rows) > 1:
            lines = ", ".join(str(row) for row in canonical_rows)
            raise ExcelStorageError(
                f"A base de alunos contém o mesmo ID ou CPF nas linhas {lines}. "
                "Revise esses cadastros antes de salvar."
            )
        canonical_id = str(ws.cell(row_idx, id_col).value or "").strip() if id_col else ""
        return row_idx, canonical_id or str(record.get("id_aluno", "") or "").strip()

    def _upsert_student(
        self, wb, record: dict[str, str], process_id: str = "", *, new_student_id: str = ""
    ) -> str:
        """Atualiza a ficha, preserva seu ID e retorna o ID usado no vínculo."""
        student_id, cpf = self._student_identity(record)
        if not student_id and not cpf and not new_student_id:
            return ""

        ws = wb[STUDENT_SHEET]
        mapping = self._header_map(ws, STUDENT_COLUMNS)
        row_idx, canonical_id = self._resolve_student_identity(ws, mapping, record)
        canonical_id = canonical_id or new_student_id or f"ALU-{self._next_student_number(wb):06d}"
        if row_idx is None:
            row_idx = ws.max_row + 1

        values = {}
        for _heading, field in STUDENT_COLUMNS:
            current = ws.cell(row_idx, mapping[field]).value if mapping.get(field) else None
            values[field] = "" if current is None else str(current)
        values["id_aluno"] = canonical_id
        for field in (
            "nome_aluno", "cpf", "rg", "data_nascimento", "data_matricula",
            "telefone", "email", "curso_origem", "ano_cursado",
        ):
            value = str(record.get(field, "") or "").strip()
            if value:
                values[field] = value
        if not values.get("importado_em"):
            values["importado_em"] = str(record.get("importado_em", "") or "").strip()
            if not values["importado_em"]:
                values["importado_em"] = datetime.now().strftime("%d/%m/%Y %H:%M")
        process_id = str(process_id or "").strip()
        process_ids = [value.strip() for value in values.get("ids_processos", "").split(";") if value.strip()]
        for value in str(record.get("ids_processos", "") or "").split(";"):
            value = value.strip()
            if value and value.casefold() not in {item.casefold() for item in process_ids}:
                process_ids.append(value)
        if process_id and process_id.casefold() not in {value.casefold() for value in process_ids}:
            process_ids.append(process_id)
        values["ids_processos"] = "; ".join(process_ids)

        for _heading, field in STUDENT_COLUMNS:
            col_idx = mapping.get(field)
            if not col_idx:
                continue
            cell = ws.cell(row_idx, col_idx)
            cell.value = str(values.get(field, "") or "")
            cell.data_type = "s"
            cell.number_format = "@"
        return canonical_id

    @staticmethod
    def _next_student_number(wb) -> int:
        pattern = re.compile(r"^ALU-(\d+)$", re.IGNORECASE)
        numbers = []
        ws = wb[STUDENT_SHEET]
        mapping = ExcelStorage._header_map(ws, STUDENT_COLUMNS)
        id_col = mapping.get("id_aluno")
        for row_idx in range(2, ws.max_row + 1):
            value = str(ws.cell(row_idx, id_col).value or "").strip() if id_col else ""
            match = pattern.match(value)
            if match:
                numbers.append(int(match.group(1)))
        return max(numbers, default=0) + 1

    def next_student_id(self) -> str:
        return f"ALU-{self._next_student_number(self._open()):06d}"

    def save_students(self, records: list[dict[str, str]]) -> list[str]:
        """Valida o lote inteiro e grava os alunos em uma única atualização da planilha."""
        records = [dict(record) for record in records if str(record.get("nome_aluno", "") or "").strip()]
        if not records:
            raise ExcelStorageError("Não encontrei alunos com nome para salvar.")
        cpfs_by_id = {}
        for record in records:
            student_id, cpf = self._student_identity(record)
            if student_id and cpf:
                if student_id in cpfs_by_id and cpfs_by_id[student_id] != cpf:
                    raise ExcelStorageError(
                        f"Conflito: o ID '{str(record.get('id_aluno', '') or '').strip()}' "
                        "aparece com CPFs diferentes no mesmo lote. Revise os dados de origem."
                    )
                cpfs_by_id[student_id] = cpf

        wb = self._open()
        next_number = self._next_student_number(wb)
        # Reserve IDs supplied in the file before allocating automatic IDs.
        for record in records:
            generated = re.fullmatch(r"ALU-(\d+)", str(record.get("id_aluno", "") or "").strip(), re.IGNORECASE)
            if generated:
                next_number = max(next_number, int(generated.group(1)) + 1)
        saved_ids = []
        saved_keys = set()
        for record in records:
            new_student_id = f"ALU-{next_number:06d}"
            student_id = self._upsert_student(wb, record, new_student_id=new_student_id)
            if student_id == new_student_id:
                next_number += 1
            if student_id.casefold() not in saved_keys:
                saved_ids.append(student_id)
                saved_keys.add(student_id.casefold())
        self._safe_save(wb, self.path)
        return saved_ids

    def save_student(self, record: dict[str, str]) -> str:
        """Grava ou atualiza uma ficha na aba Alunos e retorna seu ID."""
        return self.save_students([record])[0]

    def get_student(self, student_id: str) -> dict[str, str] | None:
        target = student_id.strip().casefold()
        return next(
            (record for record in self.list_students()
             if record.get("id_aluno", "").strip().casefold() == target),
            None,
        )

    def find_students(self, query: str = "") -> list[dict[str, str]]:
        query = query.strip().casefold()
        records = self.list_students()
        if not query:
            return records
        return [
            record for record in records
            if any(query in value.casefold() for value in record.values())
        ]

    def delete_student(self, student_id: str):
        wb = self._open()
        ws = wb[STUDENT_SHEET]
        mapping = self._header_map(ws, STUDENT_COLUMNS)
        id_col = mapping.get("id_aluno")
        process_col = mapping.get("ids_processos")
        for row_idx in range(2, ws.max_row + 1):
            value = str(ws.cell(row_idx, id_col).value or "").strip() if id_col else ""
            if value.casefold() != student_id.strip().casefold():
                continue
            links = str(ws.cell(row_idx, process_col).value or "").strip() if process_col else ""
            if links:
                raise ExcelStorageError(
                    "Este aluno está vinculado a processos. Remova ou atualize os vínculos antes de excluir."
                )
            ws.delete_rows(row_idx, 1)
            self._safe_save(wb, self.path)
            return

    def list_students(self) -> list[dict[str, str]]:
        wb = self._open()
        if STUDENT_SHEET not in wb.sheetnames:
            return []
        ws = wb[STUDENT_SHEET]
        mapping = self._header_map(ws, STUDENT_COLUMNS)
        records = []
        for row_idx in range(2, ws.max_row + 1):
            record = {}
            for _heading, field in STUDENT_COLUMNS:
                col_idx = mapping.get(field)
                value = ws.cell(row_idx, col_idx).value if col_idx else None
                record[field] = "" if value is None else str(value)
            if any(record.get(field) for field in ("id_aluno", "cpf", "nome_aluno")):
                records.append(record)
        return records

    def delete_process(self, process_id: str):
        wb = self._open()
        ws = wb[PROCESS_SHEET]
        mapping = self._header_map(ws, PROCESS_COLUMNS)
        id_col = mapping.get("id_processo")
        if id_col:
            for row_idx in range(2, ws.max_row + 1):
                value = str(ws.cell(row_idx, id_col).value or "").strip().casefold()
                if value == process_id.strip().casefold():
                    ws.delete_rows(row_idx, 1)
                    break

        eq_ws = wb[EQUIVALENCE_SHEET]
        eq_map = self._header_map(eq_ws, EQUIVALENCE_COLUMNS)
        eq_id_col = eq_map.get("id_processo")
        if eq_id_col:
            for row_idx in range(eq_ws.max_row, 1, -1):
                value = str(eq_ws.cell(row_idx, eq_id_col).value or "").strip().casefold()
                if value == process_id.strip().casefold():
                    eq_ws.delete_rows(row_idx, 1)
        self._safe_save(wb, self.path)

    def list_equivalences(self, process_id: str) -> list[dict[str, str]]:
        wb = self._open()
        ws = wb[EQUIVALENCE_SHEET]
        mapping = self._header_map(ws, EQUIVALENCE_COLUMNS)
        rows = []
        for row_idx in range(2, ws.max_row + 1):
            id_col = mapping.get("id_processo")
            value = ws.cell(row_idx, id_col).value if id_col else None
            if str(value or "").strip().casefold() != process_id.strip().casefold():
                continue
            row = {}
            for _heading, field in EQUIVALENCE_COLUMNS:
                col_idx = mapping.get(field)
                cell_value = ws.cell(row_idx, col_idx).value if col_idx else None
                row[field] = "" if cell_value is None else str(cell_value)
            rows.append(row)
        return rows

    def save_equivalences(self, process_id: str, rows: list[dict[str, str]]):
        wb = self._open()
        ws = wb[EQUIVALENCE_SHEET]
        mapping = self._header_map(ws, EQUIVALENCE_COLUMNS)
        id_col = mapping.get("id_processo")
        if id_col:
            for row_idx in range(ws.max_row, 1, -1):
                value = str(ws.cell(row_idx, id_col).value or "").strip().casefold()
                if value == process_id.strip().casefold():
                    ws.delete_rows(row_idx, 1)
        for row in rows:
            row_idx = ws.max_row + 1
            for _heading, field in EQUIVALENCE_COLUMNS:
                col_idx = mapping.get(field)
                if not col_idx:
                    continue
                value = process_id if field == "id_processo" else row.get(field, "")
                cell = ws.cell(row_idx, col_idx)
                cell.value = str(value or "")
                cell.data_type = "s"
                cell.number_format = "@"
        self._safe_save(wb, self.path)

    @staticmethod
    def _password_hash(password: str, salt: bytes) -> bytes:
        return hashlib.pbkdf2_hmac(
            "sha256", password.encode("utf-8"), salt, PASSWORD_HASH_ITERATIONS
        )

    def _user_records(self, include_credentials: bool = False) -> list[dict[str, str]]:
        wb = self._open()
        ws = wb[USER_SHEET]
        mapping = self._header_map(ws, USER_COLUMNS)
        records = []
        for row_idx in range(2, ws.max_row + 1):
            record = {}
            for field in ("username", "function", "salt", "password_hash", "must_change_password"):
                col_idx = mapping.get(field)
                value = ws.cell(row_idx, col_idx).value if col_idx else None
                record[field] = "" if value is None else str(value)
            if not record["must_change_password"]:
                # Accounts from older workbook versions did not have this flag;
                # make them rotate their password once after upgrading.
                record["must_change_password"] = "1"
            if not record["username"]:
                continue
            if include_credentials:
                records.append(record)
            else:
                records.append({"username": record["username"], "function": record["function"]})
        return records

    def list_users(self) -> list[dict[str, str]]:
        """Lista os nomes e funções cadastrados, sem expor os hashes."""
        return self._user_records()

    def create_user(
        self, username: str, password: str, function: str, must_change_password: bool = True
    ):
        username = username.strip()
        function = function.strip()
        if not username:
            raise ExcelStorageError("Informe o nome de usuário.")
        if not function:
            raise ExcelStorageError("Informe a função do usuário.")
        if len(password) < 8:
            raise ExcelStorageError("A senha precisa ter pelo menos 8 caracteres.")

        wb = self._open()
        ws = wb[USER_SHEET]
        mapping = self._header_map(ws, USER_COLUMNS)
        username_col = mapping["username"]
        for row_idx in range(2, ws.max_row + 1):
            current = str(ws.cell(row_idx, username_col).value or "").strip()
            if current.casefold() == username.casefold():
                raise ExcelStorageError("Esse nome de usuário já está cadastrado.")

        salt = os.urandom(16)
        row_idx = ws.max_row + 1
        values = {
            "username": username,
            "function": function,
            "salt": salt.hex(),
            "password_hash": self._password_hash(password, salt).hex(),
            "must_change_password": "1" if must_change_password else "0",
        }
        for _heading, field in USER_COLUMNS:
            cell = ws.cell(row_idx, mapping[field])
            cell.value = values[field]
            cell.data_type = "s"
            cell.number_format = "@"
        self._safe_save(wb, self.path)

    def authenticate_user(self, username: str, password: str) -> dict[str, str] | None:
        """Valida a senha e retorna somente a identidade pública da conta."""
        target = username.strip().casefold()
        for record in self._user_records(include_credentials=True):
            if record["username"].strip().casefold() != target:
                continue
            try:
                salt = bytes.fromhex(record["salt"])
                expected = bytes.fromhex(record["password_hash"])
            except ValueError:
                return None
            actual = self._password_hash(password, salt)
            if hmac.compare_digest(actual, expected):
                return {
                    "username": record["username"],
                    "function": record["function"],
                    "must_change_password": record["must_change_password"] == "1",
                }
            return None
        return None

    def change_user_password(self, username: str, current_password: str, new_password: str):
        if len(new_password) < 8:
            raise ExcelStorageError("A nova senha precisa ter pelo menos 8 caracteres.")

        wb = self._open()
        ws = wb[USER_SHEET]
        mapping = self._header_map(ws, USER_COLUMNS)
        username_col = mapping["username"]
        target = username.strip().casefold()
        row_idx = None
        for candidate in range(2, ws.max_row + 1):
            current = str(ws.cell(candidate, username_col).value or "").strip()
            if current.casefold() == target:
                row_idx = candidate
                break
        if row_idx is None:
            raise ExcelStorageError("Não encontrei esse usuário na planilha.")

        salt_col = mapping["salt"]
        hash_col = mapping["password_hash"]
        try:
            old_salt = bytes.fromhex(str(ws.cell(row_idx, salt_col).value or ""))
            expected = bytes.fromhex(str(ws.cell(row_idx, hash_col).value or ""))
        except ValueError as exc:
            raise ExcelStorageError("Os dados da senha estão inválidos na planilha.") from exc
        actual = self._password_hash(current_password, old_salt)
        if not hmac.compare_digest(actual, expected):
            raise ExcelStorageError("A senha atual está incorreta.")

        new_salt = os.urandom(16)
        updates = {
            "salt": new_salt.hex(),
            "password_hash": self._password_hash(new_password, new_salt).hex(),
            "must_change_password": "0",
        }
        for field, value in updates.items():
            cell = ws.cell(row_idx, mapping[field])
            cell.value = value
            cell.data_type = "s"
            cell.number_format = "@"
        self._safe_save(wb, self.path)

    def copy_users_from(self, source: "ExcelStorage"):
        """Copia as contas (incluindo hashes, nunca senhas em texto) para outra planilha."""
        if self.path.resolve() == source.path.resolve():
            return
        source_records = source._user_records(include_credentials=True)
        if not source_records:
            return

        wb = self._open()
        ws = wb[USER_SHEET]
        mapping = self._header_map(ws, USER_COLUMNS)
        username_col = mapping["username"]
        rows_by_username = {}
        for row_idx in range(2, ws.max_row + 1):
            current = str(ws.cell(row_idx, username_col).value or "").strip().casefold()
            if current:
                rows_by_username[current] = row_idx

        for record in source_records:
            key = record["username"].strip().casefold()
            row_idx = rows_by_username.get(key)
            if row_idx is None:
                row_idx = ws.max_row + 1
                rows_by_username[key] = row_idx
            for _heading, field in USER_COLUMNS:
                cell = ws.cell(row_idx, mapping[field])
                cell.value = record[field]
                cell.data_type = "s"
                cell.number_format = "@"
        self._safe_save(wb, self.path)


class StudentStorage(ExcelStorage):
    """Pasta de trabalho independente que contém a base de alunos."""

    def _open(self):
        openpyxl, Alignment, Font, PatternFill = self._load_openpyxl()
        self.path.parent.mkdir(parents=True, exist_ok=True)
        if self.path.exists():
            try:
                wb = openpyxl.load_workbook(self.path)
            except PermissionError as exc:
                raise ExcelStorageError(
                    "N\u00e3o consegui abrir a base de alunos. Feche o arquivo no Excel e tente novamente."
                ) from exc
            except Exception as exc:
                raise ExcelStorageError(
                    f"N\u00e3o consegui ler a base de alunos:\n{exc}"
                ) from exc
        else:
            wb = openpyxl.Workbook()
            wb.active.title = STUDENT_SHEET

        self._ensure_sheet(wb, STUDENT_SHEET, STUDENT_COLUMNS, Font, PatternFill, Alignment)
        return wb

    def resolve_student_id(self, record: dict[str, str]) -> str:
        """Valida a identidade e devolve o ID canônico sem gravar a base de alunos."""
        wb = self._open()
        ws = wb[STUDENT_SHEET]
        mapping = self._header_map(ws, STUDENT_COLUMNS)
        _row_idx, student_id = self._resolve_student_identity(ws, mapping, record)
        if not student_id and self._student_identity(record)[1]:
            student_id = f"ALU-{self._next_student_number(wb):06d}"
        return student_id

    def save_process_student(self, record: dict[str, str], process_id: str) -> str:
        wb = self._open()
        student_id = self._upsert_student(wb, record, process_id=process_id)
        self._safe_save(wb, self.path)
        return student_id

    def remove_process_link(self, student_id: str, process_id: str):
        wb = self._open()
        self._remove_process_link(wb, student_id, process_id)
        self._safe_save(wb, self.path)

    def unlink_process(self, process_id: str):
        wb = self._open()
        ws = wb[STUDENT_SHEET]
        mapping = self._header_map(ws, STUDENT_COLUMNS)
        links_col = mapping.get("ids_processos")
        if not links_col:
            return
        for row_idx in range(2, ws.max_row + 1):
            links = [
                value.strip()
                for value in str(ws.cell(row_idx, links_col).value or "").split(";")
                if value.strip() and value.strip().casefold() != process_id.strip().casefold()
            ]
            ws.cell(row_idx, links_col).value = "; ".join(links)
        self._safe_save(wb, self.path)
